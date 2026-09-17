"""החלת תיקונים על docx שכבר מנוקד — בלי להריץ את המודל מחדש.

ניקוד מלא של ספר לוקח כ-19 דקות על CPU, כי המודל רץ על כל פסקה. אבל
תיקון של כמה מילים לא דורש את המודל בכלל: הצורה כבר הוכרעה, וצריך רק
להחליף אותה. הכלי הזה עושה בדיוק את זה, ורץ בשניות.

    python pipeline/apply_fixes.py in.docx out.docx

הוא קורא את `lexicon/overrides.json` ו-`lexicon/homographs.json`, ומחיל
אותם על הטקסט המנוקד. זו לולאת התיקונים של שלב 6, בגרסה שאפשר לחיות
איתה: משנים כלל, מריצים, בודקים — בלי להמתין.

הקו האדום בתוקף: כל צורה נבדקת מול שלד המילה לפני שהיא נכנסת, והפסקה
נבדקת שוב אחרי הכתיבה.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from docx_pipeline import all_runs, redistribute_segments  # noqa: E402
from engine import _resolve_homograph  # noqa: E402
from hebrew import (  # noqa: E402
    identity_diff,
    is_acronym,
    is_hebrew,
    normalize,
    split_tokens,
    strip_nikud,
)

LEXICON = ROOT / "lexicon"


def _load(path: Path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def fix_paragraph(text: str, homographs: dict, overrides: dict,
                  paragraph_id: str | None = None,
                  frozen: set[int] | None = None
                  ) -> tuple[str, list[tuple[str, str]], list[tuple[str, str]]]:
    """מחזיר (טקסט מתוקן, מקטעים, שינויים)."""
    source = normalize(text)
    parts = split_tokens(source)
    word_idx = [i for i, (_t, w) in enumerate(parts) if w]
    forms = [parts[i][0] for i in word_idx]
    skeletons = [strip_nikud(f) for f in forms]

    out = [p for p, _w in parts]
    changes: list[tuple[str, str]] = []

    para_over = overrides.get(paragraph_id, {}) if paragraph_id else {}
    global_over = overrides.get("*", {})

    frozen = frozen or set()
    for i, part_i in enumerate(word_idx):
        if is_acronym(forms[i]) or i in frozen:
            continue
        skel = skeletons[i]
        new = _resolve_homograph(homographs.get(skel), skeletons, i)
        new = para_over.get(skel) or global_over.get(skel) or new
        if not new or new == forms[i]:
            continue
        # שלד שונה נפסל — תיקון משנה ניקוד, לא אותיות.
        if strip_nikud(new) != skel:
            continue
        out[part_i] = new
        changes.append((forms[i], new))

    segments = [(src, out[i]) for i, (src, _w) in enumerate(parts)]
    return "".join(out), segments, changes


def process(in_path: Path, out_path: Path, *, lexicon_dir: Path = LEXICON,
            source: Path | None = None) -> dict:
    import docx

    homographs = _load(lexicon_dir / "homographs.json", {})
    overrides = _load(lexicon_dir / "overrides.json", {})

    # ניקוד שהמחבר כתב בעצמו קפוא, בדיוק כמו בצנרת המלאה. בלי המקור
    # אי אפשר לדעת מה היה מנוקד מלכתחילה, והכלי היה דורס את המחבר.
    author_vocalized: set[tuple[int, int]] = set()
    if source is not None:
        from hebrew import has_nikud, words as _words

        src_doc = docx.Document(str(source))
        for idx, para in enumerate(src_doc.paragraphs):
            for j, w in enumerate(_words(normalize(para.text))):
                if has_nikud(w):
                    author_vocalized.add((idx, j))

    doc = docx.Document(str(in_path))
    stats: Counter = Counter()
    changed_words: Counter = Counter()
    failures: list[dict] = []

    for idx, para in enumerate(doc.paragraphs):
        text = para.text
        if not text.strip() or not is_hebrew(text):
            continue
        stats["paragraphs"] += 1

        frozen = {j for (i, j) in author_vocalized if i == idx}
        fixed, segments, changes = fix_paragraph(
            text, homographs, overrides, paragraph_id=f"p{idx}", frozen=frozen
        )
        if not changes:
            continue

        # התיקון אינו נוגע באותיות, ולכן בדיקת הזהות המקורית חלה.
        if identity_diff(text, fixed):
            failures.append({"paragraph": idx, "reason": "התיקון שינה אותיות"})
            continue

        runs = all_runs(para)
        run_texts = [r.text for r in runs]
        if not runs or "".join(run_texts) != text:
            stats["paragraphs_skipped_unmappable"] += 1
            continue

        pieces = redistribute_segments(segments, run_texts)
        if "".join(pieces) != fixed:
            failures.append({"paragraph": idx, "reason": "חלוקת runs נכשלה"})
            continue
        for run, piece in zip(runs, pieces):
            run.text = piece

        if para.text != fixed:
            failures.append({"paragraph": idx, "reason": "הפסקה שנכתבה אינה מה שאושר"})
            continue

        stats["paragraphs_changed"] += 1
        for old, new in changes:
            stats["words_changed"] += 1
            changed_words[f"{old} → {new}"] += 1

    if failures:
        raise SystemExit(
            f"נכשל ב-{len(failures)} פסקאות. לא נכתב קובץ פלט.\n"
            + json.dumps(failures[:5], ensure_ascii=False, indent=2)
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {
        "input": str(in_path),
        "output": str(out_path),
        "paragraphs": stats["paragraphs"],
        "paragraphs_changed": stats["paragraphs_changed"],
        "paragraphs_skipped_unmappable": stats["paragraphs_skipped_unmappable"],
        "words_changed": stats["words_changed"],
        "changes": changed_words.most_common(60),
    }


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="החלת תיקונים על docx מנוקד")
    p.add_argument("input", type=Path)
    p.add_argument("output", type=Path)
    p.add_argument("--lexicon-dir", type=Path, default=LEXICON)
    p.add_argument("--source", type=Path, default=None,
                   help="הספר המקורי — כדי לא לדרוס ניקוד שהמחבר כתב בעצמו")
    args = p.parse_args(argv)
    print(json.dumps(process(args.input, args.output, lexicon_dir=args.lexicon_dir,
                             source=args.source),
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

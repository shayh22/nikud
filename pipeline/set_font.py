"""קביעת הגופן של הטקסט העברי ב-docx.

למה זה נחוץ: מיקום הניקוד הוא תכונה של הגופן, לא של הטקסט. גופנים
שממקמים סימנים ביחס לצורת האות שמים את החיריק של יו"ד גבוה — היו"ד
קטנה ויושבת בראש השורה — והוא אינו מיושר עם הניקוד של שאר האותיות.
אין דרך לתקן את זה בטקסט; רק גופן שממקם ניקוד ביחס לקו הבסיס פותר זאת.

עברית ב-Word היא "כתב מורכב" (complex script), ולכן התכונה שקובעת
בפועל היא `w:cs` ולא `w:ascii`. כלי שמשנה רק את שם הגופן הרגיל לא
ישפיע על הטקסט העברי כלל.

גודל הגופן, הבולט, הנטוי, הצבע והיישור — כולם נשארים כפי שהם.

    python pipeline/set_font.py in.docx out.docx --font David
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from hebrew import is_hebrew  # noqa: E402

# גופנים שממקמים ניקוד ביחס לקו הבסיס, ולכן החיריק של יו"ד יושב בשורה
# אחת עם שאר הניקוד.
RECOMMENDED = ("David", "Narkisim", "Frank Ruehl CLM", "Taamey Frank CLM")


def set_run_font(run, name: str, *, latin_too: bool = False) -> bool:
    """קובע את הגופן של run אחד. מחזיר True אם השתנה משהו.

    כברירת מחדל נקבע רק `w:cs` — התכונה שבה Word משתמש לעברית. טקסט
    לטיני שיושב באותו run שומר על הגופן שלו, והעיצוב של המסמך אינו
    משתנה מעבר לנדרש.
    """
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        # ב-OOXML ל-rPr יש סדר ילדים מחייב: w:rFonts בא מיד אחרי
        # w:rStyle. הכנסה למקום הלא נכון פוסלת את הקובץ.
        style = rpr.find(qn("w:rStyle"))
        rpr.insert(1 if style is not None else 0, fonts)

    changed = False
    attrs = ("w:ascii", "w:hAnsi", "w:cs") if latin_too else ("w:cs",)
    for attr in attrs:
        if fonts.get(qn(attr)) != name:
            fonts.set(qn(attr), name)
            changed = True
    return changed


def apply(in_path: Path, out_path: Path, font: str, *,
          hebrew_only: bool = True, latin_too: bool = False) -> dict:
    import docx
    from docx.text.run import Run

    doc = docx.Document(str(in_path))
    stats: Counter = Counter()
    before: Counter = Counter()

    def walk(paragraphs):
        for para in paragraphs:
            for r in para._p.xpath(".//w:r"):
                run = Run(r, para)
                text = run.text
                if hebrew_only and not is_hebrew(text):
                    stats["skipped_non_hebrew"] += 1
                    continue
                # w:rFonts יושב תחת w:rPr, לא תחת w:r. חיפוש ישיר על
                # w:r מחזיר None תמיד, ואז הדיווח "אין גופן" שקרי.
                from docx.oxml.ns import qn
                rpr = r.find(qn("w:rPr"))
                fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
                before[(fonts.get(qn("w:cs")) if fonts is not None else None)
                       or "(לא נקבע)"] += 1
                did = set_run_font(run, font, latin_too=latin_too)
                stats["changed" if did else "already"] += 1

    walk(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            walk(part.paragraphs)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {
        "font": font,
        "runs_changed": stats["changed"],
        "runs_already_set": stats["already"],
        "runs_skipped_non_hebrew": stats["skipped_non_hebrew"],
        "scope": "עברית בלבד (w:cs)" if not latin_too else "עברית ולטינית",
        "previous_hebrew_fonts": dict(before.most_common(10)),
    }


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="קביעת גופן לטקסט העברי ב-docx")
    p.add_argument("input", type=Path)
    p.add_argument("output", type=Path)
    p.add_argument("--font", default="David",
                   help="מומלצים: " + ", ".join(RECOMMENDED))
    p.add_argument("--all-runs", action="store_true",
                   help="גם runs שאין בהם עברית")
    p.add_argument("--latin-too", action="store_true",
                   help="לקבוע גם את הגופן הלטיני, לא רק את זה של העברית")
    args = p.parse_args(argv)

    import json

    result = apply(args.input, args.output, args.font,
                   hebrew_only=not args.all_runs, latin_too=args.latin_too)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""בדיקות. הקו האדום נבדק כאן ראשון ובכל נתיב.

הרצה:  python -m pytest tests/ -q     או     python tests/test_engine.py
"""

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import clean  # noqa: E402
import hebrew as h  # noqa: E402
from build_lexicon import Lexicon  # noqa: E402
from engine import Engine  # noqa: E402
from quotes import QuoteIndex  # noqa: E402

VERSE = "בְּרֵאשִׁית בָּרָא אֱלֹהִים אֵת הַשָּׁמַיִם וְאֵת הָאָרֶץ"
MISHNAH = "מֵאֵימָתַי קוֹרִין אֶת שְׁמַע בְּעַרְבִית"


class TestIdentity(unittest.TestCase):
    """הקו האדום: אחרי הסרת ניקוד הטקסט חייב להיות זהה למקור תו־בתו."""

    def test_strip_is_inverse_of_vocalization(self):
        self.assertEqual(h.strip_nikud(VERSE), "בראשית ברא אלהים את השמים ואת הארץ")

    def test_assert_identity_accepts_matching(self):
        plain = h.strip_nikud(VERSE)
        self.assertEqual(h.assert_identity(plain, VERSE), VERSE)

    def test_assert_identity_rejects_eaten_letter(self):
        plain = h.strip_nikud(VERSE)
        with self.assertRaises(h.IdentityError):
            h.assert_identity(plain, VERSE.replace("אֵת", "אֵ", 1))

    def test_assert_identity_rejects_added_letter(self):
        plain = h.strip_nikud(VERSE)
        with self.assertRaises(h.IdentityError):
            h.assert_identity(plain, VERSE + "וֹ")

    def test_punctuation_is_not_nikud(self):
        """מקף, סוף פסוק ופסק הם תווים לכל דבר. הסרת ניקוד לא נוגעת בהם."""
        text = "כׇּל־הַנְּשָׁמָה תְּהַלֵּל יָהּ׃"
        self.assertEqual(h.strip_nikud(text), "כל־הנשמה תהלל יה׃")

    def test_cantillation_is_stripped(self):
        text = "בְּרֵאשִׁ֖ית בָּרָ֣א אֱלֹהִ֑ים"
        self.assertEqual(h.strip_nikud(text), "בראשית ברא אלהים")


class TestComparison(unittest.TestCase):
    def test_mark_order_does_not_affect_equality(self):
        """הסימנים נבנים בקוד: מחרוזת בקובץ מקור עוברת NFC ומאבדת את הסדר."""
        dagesh_first = "\u05d1\u05bc\u05b0"   # בית, דגש, שווא
        sheva_first = "\u05d1\u05b0\u05bc"    # בית, שווא, דגש
        self.assertNotEqual(dagesh_first, sheva_first)
        self.assertTrue(h.word_matches(dagesh_first, sheva_first))
        self.assertEqual(h.canonical(dagesh_first), h.canonical(sheva_first))

    def test_shin_dot_counts_as_an_error(self):
        self.assertFalse(h.word_matches("שָׂם", "שָׁם"))

    def test_ignore_shin_isolates_the_problem(self):
        self.assertTrue(h.word_matches("שָׂם", "שָׁם", ignore_shin=True))

    def test_char_errors_counts_slots(self):
        errors, total = h.char_errors("שָׁלוֹם", "שָׁלוּם")
        self.assertEqual(errors, 1)
        self.assertEqual(total, 4)  # ש ל ו ם

    def test_acronym_detection(self):
        self.assertTrue(h.is_acronym('רמב"ם'))
        self.assertTrue(h.is_acronym("ע'"))
        self.assertFalse(h.is_acronym("שלום"))


class TestTokenization(unittest.TestCase):
    def test_split_tokens_reassembles_exactly(self):
        text = 'אָמַר רַבִּי עֲקִיבָא: "הַכֹּל צָפוּי" — וְעוֹד.'
        self.assertEqual("".join(t for t, _ in h.split_tokens(text)), text)

    def test_words_skips_punctuation(self):
        self.assertEqual(h.words("אָב, בֵּן; רוּחַ"), ["אָב", "בֵּן", "רוּחַ"])


class TestEngineLayers(unittest.TestCase):
    def make_engine(self, **kwargs):
        lex = Lexicon(
            {
                "meta": {},
                "forms": {
                    "שלום": {"category": "unambiguous", "best": "שָׁלוֹם",
                             "count": 50, "ratio": 1.0, "variants": []},
                    "ספר": {"category": "context", "best": "סֵפֶר",
                            "count": 20, "ratio": 0.6, "variants": []},
                    "זכר": {"category": "ambiguous", "best": "זָכָר",
                            "count": 10, "ratio": 0.5, "variants": []},
                },
            },
            {"בבית\tספר": "סֵפֶר", "אמר\tספר": "סָפַר"},
            set(),
        )
        return Engine(lexicon=lex, **kwargs)

    def test_unambiguous_lexicon_applies(self):
        eng = self.make_engine()
        res = eng.vocalize("שלום")
        self.assertEqual(res.text, "שָׁלוֹם")
        self.assertEqual(res.decisions[0].source, "lexicon")

    def test_bigram_resolves_context(self):
        eng = self.make_engine()
        self.assertEqual(eng.vocalize("אמר ספר").text.split()[1], "סָפַר")
        self.assertEqual(eng.vocalize("בבית ספר").text.split()[1], "סֵפֶר")

    def test_ambiguous_is_flagged_not_guessed(self):
        """לא להמציא ניקוד. צורה דו־משמעית מסומנת לבדיקה."""
        eng = self.make_engine()
        res = eng.vocalize("זכר")
        self.assertEqual(res.text, "זכר")
        self.assertEqual(res.unresolved, ["זכר"])
        self.assertFalse(res.decisions[0].confident)

    def test_acronyms_are_never_vocalized(self):
        eng = self.make_engine()
        res = eng.vocalize('שלום רמב"ם')
        self.assertIn('רמב"ם', res.text)
        self.assertEqual(res.decisions[1].source, "acronym")

    def test_existing_nikud_is_preserved(self):
        eng = self.make_engine()
        res = eng.vocalize("שָׁלֹם")
        self.assertEqual(res.text, "שָׁלֹם")
        self.assertEqual(res.decisions[0].source, "preexisting")

    def test_override_beats_every_other_layer(self):
        eng = self.make_engine(overrides={"*": {"שלום": "שַׁלּוּם"}})
        res = eng.vocalize("שלום")
        self.assertEqual(res.text, "שַׁלּוּם")
        self.assertEqual(res.decisions[0].source, "override")

    def test_homograph_rule_beats_lexicon(self):
        eng = self.make_engine(
            homographs={"שלום": {"default": "שִׁלּוּם", "after": {"על": "שָׁלוֹם"}}}
        )
        self.assertEqual(eng.vocalize("שלום").text, "שִׁלּוּם")
        self.assertTrue(eng.vocalize("על שלום").text.endswith("שָׁלוֹם"))

    def test_layer_returning_wrong_skeleton_is_rejected(self):
        """שכבה שמחזירה צורה ששלדה שונה נפסלת — היא לא מגיעה לפלט."""
        eng = self.make_engine(overrides={"*": {"שלום": "שָׁלוֹמוֹת"}})
        res = eng.vocalize("שלום")
        self.assertEqual(res.text, "שלום")
        self.assertEqual(res.unresolved, ["שלום"])

    def test_identity_holds_on_mixed_text(self):
        eng = self.make_engine()
        text = 'שלום, זכר — רמב"ם (ספר) 1234 hello!'
        self.assertFalse(h.identity_diff(text, eng.nikud(text)))


class TestQuotes(unittest.TestCase):
    def build_index(self):
        idx = QuoteIndex(
            tokens=h.words(VERSE) + ["||"] + h.words(MISHNAH),
            index={},
            spans=[[0, 7, "Genesis 1:1"], [8, 13, "Mishnah Berakhot 1:1"]],
        )
        skel = [h.strip_nikud(t) for t in idx.tokens]
        idx._skeletons = skel
        for i in range(len(skel) - 3):
            if "||" in skel[i : i + 4]:
                continue
            idx.index[" ".join(skel[i : i + 4])] = i
        return idx

    def test_quote_is_copied_from_source(self):
        eng = Engine(quotes=self.build_index())
        res = eng.vocalize("בראשית ברא אלהים את השמים")
        self.assertEqual(res.text, "בְּרֵאשִׁית בָּרָא אֱלֹהִים אֵת הַשָּׁמַיִם")
        self.assertTrue(all(d.source == "quote" for d in res.decisions))

    def test_short_fragment_is_not_matched(self):
        eng = Engine(quotes=self.build_index())
        res = eng.vocalize("ברא אלהים")
        self.assertEqual(res.text, "ברא אלהים")

    def test_match_does_not_cross_source_boundary(self):
        eng = Engine(quotes=self.build_index())
        res = eng.vocalize("את השמים ואת הארץ מאימתי קורין את שמע")
        self.assertNotIn("||", res.text)
        self.assertFalse(h.identity_diff("את השמים ואת הארץ מאימתי קורין את שמע", res.text))


class TestClean(unittest.TestCase):
    def test_html_and_markers_removed(self):
        self.assertEqual(clean.clean_line("<b>שָׁלוֹם</b> {הערה} [1] עוֹלָם"),
                         "שָׁלוֹם עוֹלָם")

    def test_partial_nikud_is_filtered(self):
        self.assertTrue(clean.is_fully_vocalized(MISHNAH))
        self.assertFalse(clean.is_fully_vocalized("מאימתי קורין את שְׁמַע בערבית"))

    def test_sentences_split_on_sof_pasuq(self):
        got = clean.split_sentences("בְּרֵאשִׁית בָּרָא אֱלֹהִים׃ אֵת הַשָּׁמַיִם וְאֵת הָאָרֶץ׃")
        self.assertEqual(len(got), 2)


class TestDictaBertLabels(unittest.TestCase):
    """תוויות האימון חייבות להיות הפיכות — אחרת מאמנים על רעש."""

    NIKUD = ['', '<MAT_LECT>', 'ּ', 'ְ', 'ֱ', 'ֲ', 'ֳ',
             'ִ', 'ֵ', 'ֶ', 'ַ', 'ָ', 'ֹ', 'ֺ',
             'ֻ', 'ְּ', 'ֱּ', 'ֲּ', 'ֳּ',
             'ִּ', 'ֵּ', 'ֶּ', 'ַּ',
             'ָּ', 'ֹּ', 'ֺּ', 'ֻּ',
             'ׇ', 'ׇּ']
    SHIN = ['ׁ', 'ׂ']

    def roundtrip(self, gold):
        from dictabert import apply_labels, char_labels

        _bases, n, s = char_labels(gold, self.NIKUD, self.SHIN)
        plain = h.strip_nikud(h.normalize(gold))
        self.assertEqual(len(n), len(plain))
        return apply_labels(plain, n, s, self.NIKUD, self.SHIN)

    def test_roundtrip_preserves_dagesh_and_vowel(self):
        self.assertEqual(h.canonical(self.roundtrip(MISHNAH)), h.canonical(MISHNAH))

    def test_roundtrip_preserves_sin_dot(self):
        gold = "שָׂשׂוֹן וְשִׂמְחָה"
        self.assertEqual(h.canonical(self.roundtrip(gold)), h.canonical(gold))

    def test_roundtrip_preserves_qamats_qatan(self):
        gold = "כׇּל־הַנְּשָׁמָה תְּהַלֵּל יָהּ"
        self.assertEqual(h.canonical(self.roundtrip(gold)), h.canonical(gold))

    def test_chunking_reassembles_exactly(self):
        from dictabert import split_for_model

        text = " ".join(["מילה"] * 900)
        self.assertEqual("".join(split_for_model(text, 200)), text)


class TestDocxRedistribution(unittest.TestCase):
    """חלוקת הניקוד חזרה ל-runs של docx, בלי לשבור עיצוב ובלי לאבד תו."""

    def redistribute(self, vocalized, runs):
        sys.path.insert(0, str(ROOT / "pipeline"))
        from docx_pipeline import redistribute

        return redistribute(vocalized, runs)

    def test_plain_runs(self):
        pieces = self.redistribute("אָמַר רַבִּי עֲקִיבָא", ["אמר ", "רבי ", "עקיבא"])
        self.assertEqual(pieces, ["אָמַר ", "רַבִּי ", "עֲקִיבָא"])

    def test_runs_that_already_carry_nikud(self):
        """בספר יש פסוקים מנוקדים ומוטעמים. אורך הגולמי גדול מאורך השלד."""
        runs = ["הַנּוֹתֵ֥ן ", "בַּיָּ֖ם ", "דָּ֑רֶךְ"]
        pieces = self.redistribute("".join(runs), runs)
        self.assertEqual(pieces, runs)

    def test_run_boundary_inside_a_word(self):
        pieces = self.redistribute("שָׁלוֹם", ["של", "ום"])
        self.assertEqual("".join(pieces), "שָׁלוֹם")
        self.assertEqual([h.strip_nikud(p) for p in pieces], ["של", "ום"])

    def test_empty_run_is_preserved(self):
        pieces = self.redistribute("שָׁלוֹם", ["שלום", ""])
        self.assertEqual("".join(pieces), "שָׁלוֹם")
        self.assertEqual(len(pieces), 2)


class TestDocxHyperlinkRuns(unittest.TestCase):
    """רגרסיה: טקסט בתוך w:hyperlink אינו ב-paragraph.runs.

    הבאג שהיה: הצנרת הסתמכה על paragraph.runs, לא מצאה שם את טקסט
    הקישור, ונפלה למסלול "שיטוח" שכתב את כל הפסקה ל-run הראשון —
    בזמן שה-runs שבתוך הקישור נשארו כמות שהם. התוצאה: הפסקה הופיעה
    פעמיים בקובץ הפלט.
    """

    def make_doc(self):
        import docx
        from docx.oxml.ns import qn

        doc = docx.Document()
        para = doc.add_paragraph()
        para.add_run("אמר ")
        link = para._p.makeelement(qn("w:hyperlink"), {})
        run = para._p.makeelement(qn("w:r"), {})
        t = para._p.makeelement(qn("w:t"), {})
        t.text = "רבי עקיבא"
        run.append(t)
        link.append(run)
        para._p.append(link)
        return doc, para

    def test_paragraph_runs_misses_hyperlink_text(self):
        _doc, para = self.make_doc()
        self.assertEqual(para.text, "אמר רבי עקיבא")
        self.assertNotEqual("".join(r.text for r in para.runs), para.text)

    def test_all_runs_covers_hyperlink_text(self):
        sys.path.insert(0, str(ROOT / "pipeline"))
        from docx_pipeline import all_runs

        _doc, para = self.make_doc()
        self.assertEqual("".join(r.text for r in all_runs(para)), para.text)

    def test_written_paragraph_is_not_duplicated(self):
        sys.path.insert(0, str(ROOT / "pipeline"))
        from docx_pipeline import all_runs, redistribute

        _doc, para = self.make_doc()
        source = para.text
        vocalized = "אָמַר רַבִּי עֲקִיבָא"
        runs = all_runs(para)
        pieces = redistribute(vocalized, [r.text for r in runs])
        for run, piece in zip(runs, pieces):
            run.text = piece
        self.assertEqual(para.text, vocalized)
        self.assertFalse(h.identity_diff(source, para.text))


class TestKtivHaser(unittest.TestCase):
    """שכבת הכתיב החסר — השכבה היחידה שמורידה אותיות."""

    def setUp(self):
        import ktiv
        from build_lexicon import Lexicon

        # לקסיקון מזערי, כדי שהבדיקות לא יהיו תלויות בקורפוס שנבנה.
        self.lex = Lexicon(
            {"meta": {}, "forms": {
                "חבור": {"category": "unambiguous", "best": "חִבּוּר", "count": 263,
                         "ratio": 0.99, "variants": [["חִבּוּר", 260], ["חָבוֹר", 3]]},
                "חיבור": {"category": "unambiguous", "best": "חִיבּוּר", "count": 38,
                          "ratio": 1.0, "variants": [["חִיבּוּר", 38]]},
                "כלם": {"category": "context", "best": "כֻּלָּם", "count": 260,
                        "ratio": 0.89, "variants": [["כֻּלָּם", 232], ["כֻלָּם", 28]]},
                "אש": {"category": "unambiguous", "best": "אֵשׁ", "count": 269,
                       "ratio": 1.0, "variants": [["אֵשׁ", 269]]},
                "הא": {"category": "unambiguous", "best": "הָא", "count": 5862,
                       "ratio": 0.99, "variants": [["הָא", 5819], ["הִא", 1]]},
                "תקון": {"category": "context", "best": "תַּקּוּן", "count": 55,
                         "ratio": 0.42, "variants": [["תַּקּוּן", 23], ["תִקּוּן", 17],
                                                     ["תִּקּוּן", 15]]},
            }},
            {}, set(),
        )
        self.conv = ktiv.HaserConverter(self.lex)

    def test_hiriq_yod_is_dropped(self):
        got = self.conv.convert("חִיבּוּר")
        self.assertIsNotNone(got)
        self.assertEqual(got.form, "חִבּוּר")
        self.assertEqual(got.letters, ("י",))

    def test_shuruk_vav_becomes_qubuts(self):
        got = self.conv.convert("כּוּלָּם")
        self.assertIsNotNone(got)
        self.assertEqual(got.form, "כֻּלָּם")

    def test_prefixed_word_uses_stem_evidence(self):
        """הַחִיבּוּר אינו בלקסיקון, חִבּוּר כן. הראיה מהגזע מספיקה."""
        got = self.conv.convert("הַחִיבּוּר")
        self.assertIsNotNone(got)
        self.assertEqual(got.form, "הַחִבּוּר")

    def test_different_word_is_not_dropped(self):
        """אִישׁ → אֵשׁ הוא מילה אחרת. הצירה פוסל את ההורדה."""
        self.assertIsNone(self.conv.convert("אִישׁ"))

    def test_rare_matching_variant_is_not_enough(self):
        """הָא בקמץ; הִא בחיריק מופיעה פעם אחת. פעם אחת אינה ראיה."""
        self.assertIsNone(self.conv.convert("הִיא"))

    def test_word_final_yod_is_never_dropped(self):
        self.assertIsNone(self.conv.convert("רַבִּי"))

    def test_consonantal_yod_is_never_dropped(self):
        """יו"ד שנושאת ניקוד משלה אינה אם קריאה."""
        self.assertIsNone(self.conv.convert("חַיִּים"))

    def test_dagesh_is_not_lost_to_an_edition_variant(self):
        """תִקּוּן×17 שכיחה מ-תִּקּוּן×15, אבל הן אותה צורה. הדגש נשמר."""
        got = self.conv.convert("תִּיקּוּן")
        self.assertIsNotNone(got)
        self.assertEqual(got.form, "תִּקּוּן")

    def test_conversion_is_restorable(self):
        """הערובה שמחליפה את בדיקת הזהות: ההורדה הפיכה."""
        import ktiv

        got = self.conv.convert("הַחִיבּוּר")
        rebuilt = ktiv.restore(h.strip_nikud(got.form), got.removed, got.letters)
        self.assertEqual(rebuilt, h.strip_nikud("הַחִיבּוּר"))

    def test_unrestorable_conversion_is_rejected(self):
        import ktiv

        bogus = ktiv.Conversion("חִבּוּר", (9,), ("י",), 100)
        with self.assertRaises(ktiv.RestorationError):
            ktiv.assert_restorable("חִיבּוּר", bogus)


class TestEngineKtivMode(unittest.TestCase):
    def make(self, ktiv_mode):
        from build_lexicon import Lexicon

        lex = Lexicon(
            {"meta": {}, "forms": {
                "חיבור": {"category": "unambiguous", "best": "חִיבּוּר", "count": 38,
                          "ratio": 1.0, "variants": [["חִיבּוּר", 38]]},
                "חבור": {"category": "unambiguous", "best": "חִבּוּר", "count": 263,
                         "ratio": 0.99, "variants": [["חִבּוּר", 260]]},
            }},
            {}, set(),
        )
        return Engine(lexicon=lex, ktiv=ktiv_mode)

    def test_male_mode_preserves_identity(self):
        eng = self.make("male")
        res = eng.vocalize("חיבור")
        self.assertEqual(res.text, "חִיבּוּר")
        self.assertEqual(res.removals, [])
        self.assertFalse(h.identity_diff("חיבור", res.text))

    def test_haser_mode_drops_and_logs(self):
        eng = self.make("haser")
        res = eng.vocalize("חיבור")
        self.assertEqual(res.text, "חִבּוּר")
        self.assertEqual(len(res.removals), 1)
        self.assertEqual(res.removals[0]["letters"], ["י"])
        self.assertEqual(res.removed_letters, 1)

    def test_segments_reconstruct_both_sides(self):
        eng = self.make("haser")
        res = eng.vocalize("על חיבור זה")
        self.assertEqual("".join(a for a, _b in res.segments), "על חיבור זה")
        self.assertEqual("".join(b for _a, b in res.segments), res.text)

    def test_unknown_ktiv_mode_is_rejected(self):
        with self.assertRaises(ValueError):
            self.make("something")


class TestKtivRoundTrip(unittest.TestCase):
    """to_male ו-mechanical חייבות להיות הופכיות. בלי זה המדידה חסרת ערך."""

    CASES = ["חִבּוּר", "תְּפִלָּה", "כֻּלָּם", "שֻׁלְחָן", "עִקָּר", "תִּקּוּן"]

    @staticmethod
    def inserted(gold_skeleton, male_skeleton):
        """האינדקסים שבהם to_male הוסיפה אות."""
        out, i = [], 0
        for j, ch in enumerate(male_skeleton):
            if i < len(gold_skeleton) and gold_skeleton[i] == ch:
                i += 1
            else:
                out.append(j)
        return tuple(out)

    def test_male_then_mechanical_restores_the_original(self):
        """מורידים בדיוק את מה ש-to_male הוסיפה, ומקבלים את המקור.

        לא כל אם קריאה בצורה המלאה נוספה על ידי to_male: ב-חִיבּוּר גם
        הוי"ו היא אם קריאה, אבל היא הייתה שם מלכתחילה.
        """
        import ktiv

        for gold in self.CASES:
            male = ktiv.to_male(gold)
            self.assertNotEqual(male, gold, f"{gold}: to_male לא שינה דבר")
            added = self.inserted(h.strip_nikud(gold), h.strip_nikud(male))
            self.assertTrue(added, f"{male}: לא זוהתה תוספת")
            maters = ktiv.mater_positions(male)
            for pos in added:
                self.assertIn(pos, maters, f"{male}: {pos} לא זוהה כאם קריאה")
            back = ktiv.mechanical(male, added, tuple(maters[p] for p in added))
            self.assertEqual(h.canonical(back), h.canonical(gold))

    def test_male_leaves_alone_what_should_not_change(self):
        import ktiv

        for word in ["בַּיִת", "לָרִאשׁוֹן", "וְיִגְאֹל", "רַבִּי"]:
            self.assertEqual(ktiv.to_male(word), word)


class TestDocxHaserRedistribution(unittest.TestCase):
    """חלוקה ל-runs כשאותיות יורדות — שם ספירת השלד כבר לא תקפה."""

    def redistribute_segments(self, segments, runs):
        sys.path.insert(0, str(ROOT / "pipeline"))
        from docx_pipeline import redistribute_segments

        return redistribute_segments(segments, runs)

    def test_output_is_fully_preserved(self):
        segments = [("על ", "עַל "), ("חיבור", "חִבּוּר"), (" זה", " זֶה")]
        pieces = self.redistribute_segments(segments, ["על ", "חיבור", " זה"])
        self.assertEqual("".join(pieces), "עַל חִבּוּר זֶה")

    def test_run_boundary_inside_a_shortened_word(self):
        """גבול run באמצע מילה שהתקצרה — הטקסט נשמר, העיצוב מוזז."""
        segments = [("חיבור", "חִבּוּר")]
        pieces = self.redistribute_segments(segments, ["חיב", "ור"])
        self.assertEqual("".join(pieces), "חִבּוּר")
        self.assertEqual(len(pieces), 2)

    def test_more_runs_than_segments(self):
        segments = [("שלום", "שָׁלוֹם")]
        pieces = self.redistribute_segments(segments, ["שלום", "", ""])
        self.assertEqual("".join(pieces), "שָׁלוֹם")


class TestMaterVowelPlacement(unittest.TestCase):
    """תנועה שהונחה על האות שלפני וי"ו כתובה שייכת לוי"ו עצמה.

    המודל מסמן וי"ו כאם קריאה ומניח את התנועה לפניה — נכון לכתיב חסר,
    שגוי כשהוי"ו כתובה. בספר זה יצר 897 שגיאות.
    """

    def fix(self, text):
        return h.fix_mater_vowels(text)

    def test_qubuts_becomes_shuruk(self):
        self.assertEqual(self.fix("אֲרֻוכָּה"), "אֲרוּכָּה")
        self.assertEqual(self.fix("נְקֻודָּה"), "נְקוּדָּה")

    def test_holam_becomes_holam_male(self):
        self.assertEqual(self.fix("בְּאֹופֶן"), "בְּאוֹפֶן")
        self.assertEqual(self.fix("קֹודֶם"), "קוֹדֶם")

    def test_correct_forms_are_untouched(self):
        for w in ["שָׁלוֹם", "הוּא", "מִצְוָה", "עָלָיו", "תּוֹרָה", "כֻּלָּם"]:
            self.assertEqual(self.fix(w), h.normalize(w))

    def test_consonantal_vav_is_untouched(self):
        """וי"ו שנושאת תנועה או שווא היא עיצור, ולא נוגעים בה."""
        for w in ["מִצְוָה", "עֲוֹן", "וְלֹא", "תִּקְוָה"]:
            self.assertEqual(self.fix(w), h.normalize(w))

    def test_skeleton_never_changes(self):
        for w in ["אֲרֻוכָּה", "בְּאֹופֶן", "קֹודֶם", "מְסֻויֶּמֶת", "לִנְסֹועַ"]:
            self.assertEqual(h.strip_nikud(self.fix(w)), h.strip_nikud(w))

    def test_identity_still_holds(self):
        src = h.strip_nikud("אֲרֻוכָּה בְּאֹופֶן")
        self.assertFalse(h.identity_diff(src, self.fix("אֲרֻוכָּה בְּאֹופֶן")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
